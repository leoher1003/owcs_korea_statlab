#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
apply_stage2_formulas.py
Applies all Excel formulas from OWCS Korea Stage 2 to Stage 1 and Stage 3 workbooks:
1. MATCH INFO:
   - Col 10 (MAP TYPE): Dynamic map type lookup from map name (Col 9)
   - Col 15 (MATCH_ID): =TEXT(E{r}, "yymmdd") & "_" & G{r} & "_" & H{r} & "_S" & D{r}
2. MATCH_W1 ~ MATCH_W6:
   - Col 13 (Playtime(Seconds)): =L{r}*86400
   - Col 14 (MATCH_ID): ArrayFormula =INDEX('MATCH INFO'!O:O, INT((ROW()-2)/10)+2 + offset)
3. TOTAL STATISTICS:
   - Automated SUMIF formulas for Elims, Deaths, Assists, Damage, Heal, Mitigated, Playtime across MATCH_W1~W6
   - E/D ratio: =IFERROR(D{idx}/E{idx},0)
4. Per10 STATISTICS:
   - Playtime_Min: ='TOTAL STATISTICS'!K{idx}/60
   - Per10 stats: =IFERROR('TOTAL STATISTICS'!{col}/'TOTAL STATISTICS'!K{idx}*600,0)
"""

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.styles import Alignment
from pathlib import Path

MAP_TYPE_FORMULA = '=IF(OR(I{r}="Antarctic Peninsula", I{r}="Nepal", I{r}="Lijiang Tower", I{r}="Busan", I{r}="Samoa", I{r}="Oasis", I{r}="Ilios"), "Control", IF(OR(I{r}="Route 66", I{r}="Watchpoint: Gibraltar", I{r}="Dorado", I{r}="Rialto", I{r}="Shambali Monastery", I{r}="Circuit Royal", I{r}="Junkertown", I{r}="Havana"), "Escort", IF(OR(I{r}="Neon Junction", I{r}="Numbani", I{r}="Midtown", I{r}="Blizzard World", I{r}="Eichenwalde", I{r}="King\'s Row", I{r}="Paraiso", I{r}="Hollywood"), "Hybrid", IF(OR(I{r}="New Queen Street", I{r}="Esperanca", I{r}="Colosseo", I{r}="Runasapi"), "Push", IF(OR(I{r}="New Junk City", I{r}="Suravasa", I{r}="Aatlis"), "Flashpoint", "")))))'
MATCH_ID_FORMULA = '=TEXT(E{r}, "yymmdd") & "_" & G{r} & "_" & H{r} & "_S" & D{r}'

WEEK_CONFIG = {
    1: {"rows": 290, "offset": ""},
    2: {"rows": 290, "offset": " + 29"},
    3: {"rows": 340, "offset": " + 29+29"},
    4: {"rows": 320, "offset": " + 92"},
    5: {"rows": 330, "offset": " + 124"},
    6: {"rows": 310, "offset": " + 157"}
}

def apply_formulas_to_workbook(target_path: Path):
    print(f"\nProcessing {target_path}...")
    if not target_path.exists():
        print(f"Error: {target_path} does not exist.")
        return

    wb = openpyxl.load_workbook(target_path, data_only=False)
    
    # -------------------------------------------------------------
    # 1. MATCH INFO Formulas
    # -------------------------------------------------------------
    if "MATCH INFO" in wb.sheetnames:
        mi = wb["MATCH INFO"]
        max_mi_rows = max(mi.max_row, 201)
        print(f"  - Applying MATCH INFO formulas (rows 2 to {max_mi_rows})...")
        for r in range(2, max_mi_rows + 1):
            mi.cell(r, 10).value = MAP_TYPE_FORMULA.format(r=r)
            mi.cell(r, 15).value = MATCH_ID_FORMULA.format(r=r)

    # -------------------------------------------------------------
    # 2. MATCH_W1 ~ MATCH_W6 Formulas
    # -------------------------------------------------------------
    center_align = Alignment(horizontal="center", vertical="center")
    for w, cfg in WEEK_CONFIG.items():
        sname = f"MATCH_W{w}"
        if sname in wb.sheetnames:
            ws = wb[sname]
            num_rows = cfg["rows"]
            offset_str = cfg["offset"]
            print(f"  - Applying {sname} formulas (rows 2 to {num_rows + 1})...")
            for r in range(2, num_rows + 2):
                # Col 13: Playtime(Seconds)
                c13 = ws.cell(r, 13)
                c13.value = f"=L{r}*86400"
                c13.alignment = center_align

                # Col 14: MATCH_ID ArrayFormula
                c14 = ws.cell(r, 14)
                c14.value = ArrayFormula(
                    ref=f"N{r}",
                    text=f"=INDEX('MATCH INFO'!O:O, INT((ROW()-2)/10)+2{offset_str})"
                )
                c14.alignment = center_align

    # -------------------------------------------------------------
    # 3. TOTAL STATISTICS Formulas
    # -------------------------------------------------------------
    if "TOTAL STATISTICS" in wb.sheetnames and "TEAM INFO" in wb.sheetnames:
        ts = wb["TOTAL STATISTICS"]
        ti = wb["TEAM INFO"]
        
        # Determine player rows
        player_rows = []
        for r in range(2, ts.max_row + 1):
            player = ts.cell(r, 2).value
            if player:
                player_rows.append(r)
        
        print(f"  - Verifying TOTAL STATISTICS formulas for {len(player_rows)} players...")
        for idx in player_rows:
            ts.cell(idx, 4).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$F:$F)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$F:$F)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$F:$F)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$F:$F)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$F:$F)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$F:$F)"
            ts.cell(idx, 5).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$G:$G)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$G:$G)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$F:$F)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$G:$G)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$G:$G)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$G:$G)" if idx == -1 else f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$G:$G)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$G:$G)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$G:$G)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$G:$G)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$G:$G)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$G:$G)"
            ts.cell(idx, 6).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$H:$H)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$H:$H)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$H:$H)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$H:$H)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$H:$H)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$H:$H)"
            ts.cell(idx, 7).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$I:$I)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$I:$I)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$I:$I)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$I:$I)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$I:$I)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$I:$I)"
            ts.cell(idx, 8).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$J:$J)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$J:$J)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$J:$J)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$J:$J)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$J:$J)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$J:$J)"
            ts.cell(idx, 9).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$K:$K)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$K:$K)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$K:$K)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$K:$K)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$K:$K)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$K:$K)"
            ts.cell(idx, 10).value = f"=IFERROR(D{idx}/E{idx},0)"
            ts.cell(idx, 11).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$M:$M)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$M:$M)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$M:$M)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$M:$M)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$M:$M)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$M:$M)"
            
            # Formats
            for c in range(4, 10):
                ts.cell(idx, c).number_format = "#,##0.00_ "
                ts.cell(idx, c).alignment = center_align
            ts.cell(idx, 10).number_format = "0.00_ "
            ts.cell(idx, 10).alignment = center_align
            ts.cell(idx, 11).number_format = "#,##0.00_ "
            ts.cell(idx, 11).alignment = center_align

        # Delete any excess trailing rows
        if ts.max_row > len(player_rows) + 1:
            ts.delete_rows(len(player_rows) + 2, ts.max_row - len(player_rows) - 1)

    # -------------------------------------------------------------
    # 4. Per10 STATISTICS Formulas
    # -------------------------------------------------------------
    if "Per10 STATISTICS" in wb.sheetnames and "TOTAL STATISTICS" in wb.sheetnames:
        p10 = wb["Per10 STATISTICS"]
        print(f"  - Verifying Per10 STATISTICS formulas for {len(player_rows)} players...")
        for idx in player_rows:
            p10.cell(idx, 4).value = f"='TOTAL STATISTICS'!K{idx}/60"
            p10.cell(idx, 5).value = f"=IFERROR('TOTAL STATISTICS'!D{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
            p10.cell(idx, 6).value = f"=IFERROR('TOTAL STATISTICS'!E{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
            p10.cell(idx, 7).value = f"=IFERROR('TOTAL STATISTICS'!F{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
            p10.cell(idx, 8).value = f"=IFERROR('TOTAL STATISTICS'!G{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
            p10.cell(idx, 9).value = f"=IFERROR('TOTAL STATISTICS'!H{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
            p10.cell(idx, 10).value = f"=IFERROR('TOTAL STATISTICS'!I{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
            for c in range(4, 11):
                p10.cell(idx, c).number_format = "0.00"

        # Delete any excess trailing rows
        if p10.max_row > len(player_rows) + 1:
            p10.delete_rows(len(player_rows) + 2, p10.max_row - len(player_rows) - 1)

    wb.save(target_path)
    print(f"✓ Successfully updated all Stage 2 formulas in {target_path}!")

def main():
    stage1_xlsx = Path("OWCS_STAT_LAB_26_KR_STAGE1.xlsx")
    stage3_xlsx = Path("OWCS_STAT_LAB_26_KR_STAGE3.xlsx")
    apply_formulas_to_workbook(stage1_xlsx)
    apply_formulas_to_workbook(stage3_xlsx)

if __name__ == "__main__":
    main()
