#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
create_stage1_workbook.py
Generates OWCS_STAT_LAB_26_KR_STAGE1.xlsx based on OWCS Korea Stage 2 structure,
updating TEAM INFO with 2026 Stage 1 confirmed rosters, clearing match data,
and rebuilding automated SUMIF and Per10 statistics formulas.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

SOURCE_XLSX = Path("OWCS_STAT_LAB_26_KR_STAGE2_RECALCULATED.xlsx")
TARGET_XLSX = Path("OWCS_STAT_LAB_26_KR_STAGE1.xlsx")

# OWCS Korea 2026 Stage 1 Confirmed Rosters & Teams (9 Teams)
STAGE1_TEAMS = [
    {
        "team_name": "Crazy Raccoon",
        "short": "CR",
        "players": [
            {"name": "JUNBIN", "pos": "TANK", "detailed": "TANK"},
            {"name": "MAX", "pos": "TANK", "detailed": "TANK"},
            {"name": "HEESANG", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "LIP", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "STALK3R", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "CH0R0NG", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "VIGILANTE", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    },
    {
        "team_name": "Team Falcons",
        "short": "FLC",
        "players": [
            {"name": "SOMEONE", "pos": "TANK", "detailed": "TANK"},
            {"name": "HANBIN", "pos": "TANK", "detailed": "TANK"},
            {"name": "CHECKMATE", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "MER1T", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "CHIYO", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "FIELDER", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    },
    {
        "team_name": "T1",
        "short": "T1",
        "players": [
            {"name": "DONGHAK", "pos": "TANK", "detailed": "TANK"},
            {"name": "JASM1NE", "pos": "TANK", "detailed": "TANK"},
            {"name": "ZEST", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "PROUD", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "BLISS", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "SKEWED", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    },
    {
        "team_name": "ZETA DIVISION",
        "short": "ZETA",
        "players": [
            {"name": "BERNAR", "pos": "TANK", "detailed": "TANK"},
            {"name": "MEALGARU", "pos": "TANK", "detailed": "TANK"},
            {"name": "PROPER", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "KNIFE", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "VIOL2T", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "SHU", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    },
    {
        "team_name": "New Era",
        "short": "ERA",
        "players": [
            {"name": "SOLA", "pos": "TANK", "detailed": "TANK"},
            {"name": "YATE", "pos": "TANK", "detailed": "TANK"},
            {"name": "D0DO", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "PERR", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "SECRET", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "MCD", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    },
    {
        "team_name": "ONSIDE GAMING",
        "short": "OSG",
        "players": [
            {"name": "ATTACK", "pos": "TANK", "detailed": "TANK"},
            {"name": "SP1NT", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "KILO", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "HAKSAL", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "OPENER", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "IRONY", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    },
    {
        "team_name": "ZAN Esports",
        "short": "ZAN",
        "players": [
            {"name": "HEISER", "pos": "TANK", "detailed": "TANK"},
            {"name": "A1IEN", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "PROBE", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "BECKY", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "HAVIRA", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "YANGIUN", "pos": "SPT", "detailed": "FLEX SPT"},
            {"name": "KIVIS", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    },
    {
        "team_name": "Cheeseburger",
        "short": "CB",
        "players": [
            {"name": "FARMER", "pos": "TANK", "detailed": "TANK"},
            {"name": "SEUNGAN", "pos": "TANK", "detailed": "TANK"},
            {"name": "ARGON", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "ZESIN", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "JAMELGONG", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "WOOCHAN", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "FAITH", "pos": "SPT", "detailed": "MAIN SPT"}
        ]
    },
    {
        "team_name": "Poker Face",
        "short": "PF",
        "players": [
            {"name": "FEARFUL", "pos": "TANK", "detailed": "TANK"},
            {"name": "GUR3UM", "pos": "TANK", "detailed": "TANK"},
            {"name": "M1NUT2", "pos": "DPS", "detailed": "MAIN DPS"},
            {"name": "D4RT", "pos": "DPS", "detailed": "FLEX DPS"},
            {"name": "TENTEN", "pos": "SPT", "detailed": "MAIN SPT"},
            {"name": "CARU", "pos": "SPT", "detailed": "FLEX SPT"},
            {"name": "SP1NEL", "pos": "SPT", "detailed": "FLEX SPT"}
        ]
    }
]

def build_stage1_workbook():
    print(f"1. Loading Stage 2 workbook template: {SOURCE_XLSX}")
    wb = openpyxl.load_workbook(SOURCE_XLSX)
    
    # -------------------------------------------------------------
    # 1. Update TEAM INFO Sheet
    # -------------------------------------------------------------
    print("2. Rebuilding 'TEAM INFO' sheet for Stage 1...")
    ti = wb["TEAM INFO"]
    # Unmerge any existing merged ranges
    merged_ranges = list(ti.merged_cells.ranges)
    for m_range in merged_ranges:
        ti.unmerge_cells(str(m_range))

    # Clear existing content below header
    for r in range(2, ti.max_row + 1):
        for c in range(1, ti.max_column + 1):
            cell = ti.cell(r, c)
            if type(cell).__name__ != 'MergedCell':
                cell.value = None

    row_idx = 2
    all_players = []
    for t_data in STAGE1_TEAMS:
        t_name = t_data["team_name"]
        short = t_data["short"]
        first = True
        for p in t_data["players"]:
            ti.cell(row_idx, 1).value = t_name if first else None
            ti.cell(row_idx, 2).value = short if first else None
            ti.cell(row_idx, 3).value = p["name"]
            ti.cell(row_idx, 4).value = p["pos"]
            ti.cell(row_idx, 5).value = p["detailed"]
            first = False
            all_players.append({"team": short, "player": p["name"], "pos": p["pos"], "detailed": p["detailed"]})
            row_idx += 1

    if ti.max_row >= row_idx:
        ti.delete_rows(row_idx, ti.max_row - row_idx + 1)

    # -------------------------------------------------------------
    # 2. Reset MATCH INFO (Header preserved, clear old matches and populate formulas)
    # -------------------------------------------------------------
    print("3. Resetting MATCH INFO sheet and applying Stage 2 formulas...")
    mi = wb["MATCH INFO"]
    for r in range(2, mi.max_row + 1):
        for c in range(1, mi.max_column + 1):
            mi.cell(r, c).value = None

    map_type_formula = '=IF(OR(I{r}="Antarctic Peninsula", I{r}="Nepal", I{r}="Lijiang Tower", I{r}="Busan", I{r}="Samoa", I{r}="Oasis", I{r}="Ilios"), "Control", IF(OR(I{r}="Route 66", I{r}="Watchpoint: Gibraltar", I{r}="Dorado", I{r}="Rialto", I{r}="Shambali Monastery", I{r}="Circuit Royal", I{r}="Junkertown", I{r}="Havana"), "Escort", IF(OR(I{r}="Neon Junction", I{r}="Numbani", I{r}="Midtown", I{r}="Blizzard World", I{r}="Eichenwalde", I{r}="King\'s Row", I{r}="Paraiso", I{r}="Hollywood"), "Hybrid", IF(OR(I{r}="New Queen Street", I{r}="Esperanca", I{r}="Colosseo", I{r}="Runasapi"), "Push", IF(OR(I{r}="New Junk City", I{r}="Suravasa", I{r}="Aatlis"), "Flashpoint", "")))))'
    match_id_formula = '=TEXT(E{r}, "yymmdd") & "_" & G{r} & "_" & H{r} & "_S" & D{r}'
    for r in range(2, max(mi.max_row, 201) + 1):
        mi.cell(r, 10).value = map_type_formula.format(r=r)
        mi.cell(r, 15).value = match_id_formula.format(r=r)

    # Clear MATCH_W1 through MATCH_W6 and populate formulas
    print("4. Resetting MATCH_W1 ~ MATCH_W6 match sheets with Stage 2 formulas...")
    from openpyxl.worksheet.formula import ArrayFormula
    center_align = Alignment(horizontal="center", vertical="center")
    week_cfg = {
        1: {"rows": 290, "offset": ""},
        2: {"rows": 290, "offset": " + 29"},
        3: {"rows": 340, "offset": " + 29+29"},
        4: {"rows": 320, "offset": " + 92"},
        5: {"rows": 330, "offset": " + 124"},
        6: {"rows": 310, "offset": " + 157"}
    }
    for w_idx in range(1, 7):
        w_sheet_name = f"MATCH_W{w_idx}"
        if w_sheet_name in wb.sheetnames:
            ws = wb[w_sheet_name]
            for r in range(2, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).value = None
            cfg = week_cfg.get(w_idx, {"rows": 290, "offset": ""})
            for r in range(2, cfg["rows"] + 2):
                c13 = ws.cell(r, 13)
                c13.value = f"=L{r}*86400"
                c13.alignment = center_align

                c14 = ws.cell(r, 14)
                c14.value = ArrayFormula(ref=f"N{r}", text=f"=INDEX('MATCH INFO'!O:O, INT((ROW()-2)/10)+2{cfg['offset']})")
                c14.alignment = center_align

    # -------------------------------------------------------------
    # 3. Rebuild TOTAL STATISTICS Sheet formulas
    # -------------------------------------------------------------
    print("5. Rebuilding 'TOTAL STATISTICS' formulas for Stage 1 players...")
    ts = wb["TOTAL STATISTICS"]
    for r in range(2, ts.max_row + 1):
        for c in range(1, ts.max_column + 1):
            ts.cell(r, c).value = None

    for idx, p in enumerate(all_players, start=2):
        ts.cell(idx, 1).value = p["team"]
        ts.cell(idx, 2).value = p["player"]
        ts.cell(idx, 3).value = p["pos"]
        # Formula columns (D through K)
        # TOTAL ELIMS
        ts.cell(idx, 4).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$F:$F)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$F:$F)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$F:$F)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$F:$F)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$F:$F)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$F:$F)"
        # TOTAL DEATHS
        ts.cell(idx, 5).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$G:$G)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$G:$G)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$G:$G)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$G:$G)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$G:$G)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$G:$G)"
        # TOTAL ASSISTS
        ts.cell(idx, 6).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$H:$H)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$H:$H)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$H:$H)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$H:$H)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$H:$H)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$H:$H)"
        # TOTAL DAMAGE
        ts.cell(idx, 7).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$I:$I)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$I:$I)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$I:$I)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$I:$I)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$I:$I)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$I:$I)"
        # TOTAL HEAL
        ts.cell(idx, 8).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$J:$J)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$J:$J)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$J:$J)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$J:$J)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$J:$J)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$J:$J)"
        # TOTAL MITIGATED
        ts.cell(idx, 9).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$K:$K)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$K:$K)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$K:$K)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$K:$K)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$K:$K)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$K:$K)"
        # TOTAL E/D
        ts.cell(idx, 10).value = f"=IFERROR(D{idx}/E{idx},0)"
        # TOTAL PLAYTIME
        ts.cell(idx, 11).value = f"=SUMIF(MATCH_W1!$B:$B,$B{idx},MATCH_W1!$M:$M)+SUMIF(MATCH_W2!$B:$B,$B{idx},MATCH_W2!$M:$M)+SUMIF(MATCH_W3!$B:$B,$B{idx},MATCH_W3!$M:$M)+SUMIF(MATCH_W4!$B:$B,$B{idx},MATCH_W4!$M:$M)+SUMIF(MATCH_W5!$B:$B,$B{idx},MATCH_W5!$M:$M)+SUMIF(MATCH_W6!$B:$B,$B{idx},MATCH_W6!$M:$M)"

    if ts.max_row > len(all_players) + 1:
        ts.delete_rows(len(all_players) + 2, ts.max_row - len(all_players) - 1)

    # -------------------------------------------------------------
    # 4. Rebuild Per10 STATISTICS Sheet formulas
    # -------------------------------------------------------------
    print("6. Rebuilding 'Per10 STATISTICS' formulas...")
    p10 = wb["Per10 STATISTICS"]
    for r in range(2, p10.max_row + 1):
        for c in range(1, p10.max_column + 1):
            p10.cell(r, c).value = None

    for idx, p in enumerate(all_players, start=2):
        p10.cell(idx, 1).value = p["team"]
        p10.cell(idx, 2).value = p["player"]
        p10.cell(idx, 3).value = p["pos"]
        p10.cell(idx, 4).value = f"='TOTAL STATISTICS'!K{idx}/60"
        p10.cell(idx, 5).value = f"=IFERROR('TOTAL STATISTICS'!D{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
        p10.cell(idx, 6).value = f"=IFERROR('TOTAL STATISTICS'!E{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
        p10.cell(idx, 7).value = f"=IFERROR('TOTAL STATISTICS'!F{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
        p10.cell(idx, 8).value = f"=IFERROR('TOTAL STATISTICS'!G{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
        p10.cell(idx, 9).value = f"=IFERROR('TOTAL STATISTICS'!H{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"
        p10.cell(idx, 10).value = f"=IFERROR('TOTAL STATISTICS'!I{idx}/'TOTAL STATISTICS'!K{idx}*600,0)"

    if p10.max_row > len(all_players) + 1:
        p10.delete_rows(len(all_players) + 2, p10.max_row - len(all_players) - 1)

    # -------------------------------------------------------------
    # 5. Clear POTM Sheet
    # -------------------------------------------------------------
    print("7. Clearing POTM sheet...")
    potm = wb["POTM"]
    for r in range(2, potm.max_row + 1):
        for c in range(1, potm.max_column + 1):
            potm.cell(r, c).value = None

    wb.save(TARGET_XLSX)
    print(f"✓ OWCS Korea 2026 Stage 1 워크북 생성 완료: {TARGET_XLSX.resolve()}")
    print(f"  - 총 등록 팀 수: {len(STAGE1_TEAMS)}개 팀")
    print(f"  - 총 등록 선수 수: {len(all_players)}명")

if __name__ == "__main__":
    build_stage1_workbook()
