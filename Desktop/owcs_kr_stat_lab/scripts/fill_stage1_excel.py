#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fill_stage1_excel.py
Fills OWCS_STAT_LAB_26_KR_STAGE1.xlsx with the complete 2026 Stage 1 matches
scraped from Liquipedia (51 matches, 186 sets, bans, detail scores, POTMs).
"""

import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

MAP_TYPE_MAP = {
    # Control
    "Antarctic Peninsula": "Control", "Nepal": "Control", "Lijiang Tower": "Control",
    "Busan": "Control", "Samoa": "Control", "Oasis": "Control", "Ilios": "Control",
    # Escort
    "Route 66": "Escort", "Watchpoint: Gibraltar": "Escort", "Dorado": "Escort",
    "Rialto": "Escort", "Shambali Monastery": "Escort", "Circuit Royal": "Escort",
    "Junkertown": "Escort", "Havana": "Escort",
    # Hybrid
    "Neon Junction": "Hybrid", "Numbani": "Hybrid", "Midtown": "Hybrid",
    "Blizzard World": "Hybrid", "Eichenwalde": "Hybrid", "King's Row": "Hybrid",
    "Paraíso": "Hybrid", "Paraiso": "Hybrid", "Hollywood": "Hybrid",
    # Push
    "New Queen Street": "Push", "Esperança": "Push", "Esperanca": "Push",
    "Colosseo": "Push", "Runasapi": "Push",
    # Flashpoint
    "New Junk City": "Flashpoint", "Suravasa": "Flashpoint", "Aatlis": "Flashpoint"
}

def fill_stage1_workbook():
    xlsx_path = "OWCS_STAT_LAB_26_KR_STAGE1.xlsx"
    json_path = "owcs-stat-lab 2/data/stage1_matches.json"
    
    with open(json_path, "r", encoding="utf-8") as f:
        matches = json.load(f)
        
    wb = openpyxl.load_workbook(xlsx_path)
    
    # 1. Update MATCH INFO
    ws_match = wb["MATCH INFO"]
    # Clear existing data rows (keep header row 1)
    while ws_match.max_row > 1:
        ws_match.delete_rows(2)
        
    row_idx = 2
    match_counter = 0
    
    # Player position lookup from TEAM INFO
    ws_team = wb["TEAM INFO"]
    player_pos = {}
    for r in ws_team.iter_rows(min_row=2, values_only=True):
        p_name = r[2] # Player Name
        pos = r[4] or r[3] # Detailed or Player Position
        if p_name and pos:
            player_pos[str(p_name).strip().upper()] = str(pos).strip()

    potm_rows = []

    for m in matches:
        match_counter += 1
        week = m.get("week") or (5 if m["phase"] in ["2nd RR", "LCQ"] else 6)
        # Parse date
        date_str = m.get("date", "").split(" - ")[0].strip()
        phase_str = "Round Robin" if m["phase"] == "Round Robin" else m["phase"]
        
        # Track POTM for this match
        if m.get("mvp"):
            mvp_name = m["mvp"].strip()
            pos = player_pos.get(mvp_name.upper(), "DPS")
            potm_rows.append((date_str, f"{m['team1']} vs {m['team2']}", mvp_name, pos))

        day_num = 1 # approximate if not explicitly known
        
        for s in m["sets"]:
            set_num = s["setNumber"]
            map_name = s["map"]
            map_type = MAP_TYPE_MAP.get(map_name, s.get("mode", "").capitalize())
            winner = s["winner"]
            ban1 = s["team1Ban"] if s["team1Ban"] else ""
            ban2 = s["team2Ban"] if s["team2Ban"] else ""
            
            # scores
            try:
                sc1 = float(s["score1"]) if "." in str(s["score1"]) else int(s["score1"])
            except:
                sc1 = s["score1"]
            try:
                sc2 = float(s["score2"]) if "." in str(s["score2"]) else int(s["score2"])
            except:
                sc2 = s["score2"]
                
            banstart_side = "TEAM 1" if s.get("banStart") == "1" else ("TEAM 2" if s.get("banStart") == "2" else "")
            initial_ban = ban1 if s.get("banStart") == "1" else (ban2 if s.get("banStart") == "2" else "")
            
            # MATCH_ID formula
            match_id_formula = f'=TEXT(E{row_idx}, "yymmdd") & "_" & G{row_idx} & "_" & H{row_idx} & "_S" & D{row_idx}'
            
            row_data = [
                week,                   # WEEK
                day_num,                # DAY
                m.get("matchNumber", 1),# MATCH
                set_num,                # Set #
                date_str,               # DATE
                phase_str,              # PHASE
                m["team1"],             # TEAM 1
                m["team2"],             # TEAM 2
                map_name,               # MAP
                map_type,               # MAP TYPE
                None,                   # TIME
                winner,                 # WINNER
                ban1,                   # TEAM 1 BAN
                ban2,                   # TEAM 2 BAN
                match_id_formula,       # MATCH_ID
                sc1,                    # TEAM 1 SCORE
                sc2,                    # TEAM 2 SCORE
                banstart_side,          # INITIAL BAN RIGHT
                initial_ban             # INITIAL BAN
            ]
            ws_match.append(row_data)
            row_idx += 1

    print(f"Filled MATCH INFO with {row_idx - 2} set rows across {len(matches)} matches.")
    
    # 2. Update POTM Sheet
    if "POTM" in wb.sheetnames:
        ws_potm = wb["POTM"]
        while ws_potm.max_row > 1:
            ws_potm.delete_rows(2)
        for pr in potm_rows:
            ws_potm.append(list(pr))
        print(f"Filled POTM sheet with {len(potm_rows)} awards.")
        
    wb.save(xlsx_path)
    print(f"Successfully saved updated workbook: {xlsx_path}")

if __name__ == "__main__":
    fill_stage1_workbook()
