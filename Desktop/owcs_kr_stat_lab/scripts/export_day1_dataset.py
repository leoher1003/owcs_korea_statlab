#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
OWWC 2026 Playoffs Day 1 Result Dataset Generator
Creates structured CSV and Excel files from the extracted high-resolution scoreboards.
"""

import pandas as pd
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

COLUMNS = [
    "Team", "Player", "Position", "Map", "Map Type",
    "Elim", "Death", "Assists", "Damage", "Heal", "Mitigated", "Playtime"
]

records = [
    # ==========================================
    # Match 1: KSA vs ESP (Map 1: Nepal)
    # ==========================================
    # KSA (Winner)
    {"Team": "KSA", "Player": "LBBD7", "Position": "DPS", "Map": "Nepal", "Map Type": "Control", "Elim": 20, "Death": 3, "Assists": 3, "Damage": 7236, "Heal": 0, "Mitigated": 263, "Playtime": "08:30"},
    {"Team": "KSA", "Player": "QUARTZ", "Position": "DPS", "Map": "Nepal", "Map Type": "Control", "Elim": 21, "Death": 3, "Assists": 0, "Damage": 7091, "Heal": 0, "Mitigated": 0, "Playtime": "08:30"},
    {"Team": "KSA", "Player": "SIRMAJED", "Position": "SPT", "Map": "Nepal", "Map Type": "Control", "Elim": 11, "Death": 1, "Assists": 21, "Damage": 2591, "Heal": 9224, "Mitigated": 0, "Playtime": "08:30"},
    {"Team": "KSA", "Player": "ZIYAD", "Position": "TANK", "Map": "Nepal", "Map Type": "Control", "Elim": 18, "Death": 3, "Assists": 5, "Damage": 7724, "Heal": 1523, "Mitigated": 3512, "Playtime": "08:30"},
    {"Team": "KSA", "Player": "ZOX", "Position": "SPT", "Map": "Nepal", "Map Type": "Control", "Elim": 12, "Death": 2, "Assists": 16, "Damage": 2429, "Heal": 8701, "Mitigated": 0, "Playtime": "08:30"},
    # ESP
    {"Team": "ESP", "Player": "BONES", "Position": "DPS", "Map": "Nepal", "Map Type": "Control", "Elim": 4, "Death": 9, "Assists": 0, "Damage": 4242, "Heal": 20, "Mitigated": 291, "Playtime": "08:30"},
    {"Team": "ESP", "Player": "GALAA", "Position": "SPT", "Map": "Nepal", "Map Type": "Control", "Elim": 5, "Death": 6, "Assists": 8, "Damage": 3687, "Heal": 6883, "Mitigated": 2165, "Playtime": "08:30"},
    {"Team": "ESP", "Player": "KHENAIL", "Position": "SPT", "Map": "Nepal", "Map Type": "Control", "Elim": 6, "Death": 6, "Assists": 6, "Damage": 4505, "Heal": 7240, "Mitigated": 0, "Playtime": "08:30"},
    {"Team": "ESP", "Player": "SANTY", "Position": "TANK", "Map": "Nepal", "Map Type": "Control", "Elim": 6, "Death": 5, "Assists": 2, "Damage": 8498, "Heal": 0, "Mitigated": 11700, "Playtime": "08:30"},
    {"Team": "ESP", "Player": "XZODYAL", "Position": "DPS", "Map": "Nepal", "Map Type": "Control", "Elim": 6, "Death": 6, "Assists": 0, "Damage": 7654, "Heal": 0, "Mitigated": 18, "Playtime": "08:30"},

    # ==========================================
    # Match 1: KSA vs ESP (Map 2: Route 66)
    # ==========================================
    # ESP (Winner)
    {"Team": "ESP", "Player": "BONES", "Position": "DPS", "Map": "Route 66", "Map Type": "Escort", "Elim": 17, "Death": 8, "Assists": 0, "Damage": 8851, "Heal": 0, "Mitigated": 197, "Playtime": "12:15"},
    {"Team": "ESP", "Player": "GALAA", "Position": "SPT", "Map": "Route 66", "Map Type": "Escort", "Elim": 14, "Death": 9, "Assists": 18, "Damage": 9181, "Heal": 12114, "Mitigated": 0, "Playtime": "12:15"},
    {"Team": "ESP", "Player": "KHENAIL", "Position": "SPT", "Map": "Route 66", "Map Type": "Escort", "Elim": 20, "Death": 11, "Assists": 12, "Damage": 8130, "Heal": 10620, "Mitigated": 487, "Playtime": "12:15"},
    {"Team": "ESP", "Player": "SANTY", "Position": "TANK", "Map": "Route 66", "Map Type": "Escort", "Elim": 21, "Death": 10, "Assists": 8, "Damage": 14739, "Heal": 0, "Mitigated": 16537, "Playtime": "12:15"},
    {"Team": "ESP", "Player": "XZODYAL", "Position": "DPS", "Map": "Route 66", "Map Type": "Escort", "Elim": 28, "Death": 12, "Assists": 2, "Damage": 16637, "Heal": 271, "Mitigated": 196, "Playtime": "12:15"},
    # KSA
    {"Team": "KSA", "Player": "LBBD7", "Position": "DPS", "Map": "Route 66", "Map Type": "Escort", "Elim": 23, "Death": 11, "Assists": 2, "Damage": 10517, "Heal": 150, "Mitigated": 633, "Playtime": "12:15"},
    {"Team": "KSA", "Player": "QUARTZ", "Position": "DPS", "Map": "Route 66", "Map Type": "Escort", "Elim": 27, "Death": 10, "Assists": 2, "Damage": 14385, "Heal": 0, "Mitigated": 844, "Playtime": "12:15"},
    {"Team": "KSA", "Player": "SIRMAJED", "Position": "SPT", "Map": "Route 66", "Map Type": "Escort", "Elim": 16, "Death": 6, "Assists": 25, "Damage": 4843, "Heal": 17942, "Mitigated": 0, "Playtime": "12:15"},
    {"Team": "KSA", "Player": "ZIYAD", "Position": "TANK", "Map": "Route 66", "Map Type": "Escort", "Elim": 32, "Death": 7, "Assists": 14, "Damage": 13180, "Heal": 2678, "Mitigated": 6767, "Playtime": "12:15"},
    {"Team": "KSA", "Player": "ZOX", "Position": "SPT", "Map": "Route 66", "Map Type": "Escort", "Elim": 22, "Death": 8, "Assists": 32, "Damage": 3884, "Heal": 14338, "Mitigated": 0, "Playtime": "12:15"},

    # ==========================================
    # Match 1: KSA vs ESP (Map 3: Eichenwalde)
    # ==========================================
    # KSA (Winner)
    {"Team": "KSA", "Player": "LBBD7", "Position": "DPS", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 19, "Death": 9, "Assists": 3, "Damage": 13440, "Heal": 155, "Mitigated": 922, "Playtime": "15:45"},
    {"Team": "KSA", "Player": "QUARTZ", "Position": "DPS", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 32, "Death": 8, "Assists": 3, "Damage": 21972, "Heal": 0, "Mitigated": 638, "Playtime": "15:45"},
    {"Team": "KSA", "Player": "SIRMAJED", "Position": "SPT", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 10, "Death": 7, "Assists": 30, "Damage": 4930, "Heal": 17840, "Mitigated": 0, "Playtime": "15:45"},
    {"Team": "KSA", "Player": "ZIYAD", "Position": "TANK", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 34, "Death": 6, "Assists": 14, "Damage": 13273, "Heal": 2875, "Mitigated": 17310, "Playtime": "15:45"},
    {"Team": "KSA", "Player": "ZOX", "Position": "SPT", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 20, "Death": 7, "Assists": 25, "Damage": 4269, "Heal": 14314, "Mitigated": 0, "Playtime": "15:45"},
    # ESP
    {"Team": "ESP", "Player": "BONES", "Position": "DPS", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 16, "Death": 9, "Assists": 1, "Damage": 8462, "Heal": 33, "Mitigated": 699, "Playtime": "15:45"},
    {"Team": "ESP", "Player": "GALAA", "Position": "SPT", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 10, "Death": 9, "Assists": 17, "Damage": 4482, "Heal": 16540, "Mitigated": 0, "Playtime": "15:45"},
    {"Team": "ESP", "Player": "KHENAIL", "Position": "SPT", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 8, "Death": 9, "Assists": 17, "Damage": 6702, "Heal": 15913, "Mitigated": 2920, "Playtime": "15:45"},
    {"Team": "ESP", "Player": "SANTY", "Position": "TANK", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 18, "Death": 8, "Assists": 4, "Damage": 17262, "Heal": 0, "Mitigated": 16789, "Playtime": "15:45"},
    {"Team": "ESP", "Player": "XZODYAL", "Position": "DPS", "Map": "Eichenwalde", "Map Type": "Hybrid", "Elim": 22, "Death": 13, "Assists": 6, "Damage": 18463, "Heal": 0, "Mitigated": 696, "Playtime": "15:45"},

    # ==========================================
    # Match 3: FRA vs AUS (Map 1: Samoa)
    # ==========================================
    # FRA (Winner)
    {"Team": "FRA", "Player": "DIP", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 17, "Death": 1, "Assists": 0, "Damage": 4708, "Heal": 87, "Mitigated": 0, "Playtime": "08:15"},
    {"Team": "FRA", "Player": "FDGOD", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 11, "Death": 2, "Assists": 13, "Damage": 1407, "Heal": 5568, "Mitigated": 0, "Playtime": "08:15"},
    {"Team": "FRA", "Player": "KIO", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 18, "Death": 3, "Assists": 1, "Damage": 6317, "Heal": 0, "Mitigated": 296, "Playtime": "08:15"},
    {"Team": "FRA", "Player": "KROXZ", "Position": "TANK", "Map": "Samoa", "Map Type": "Control", "Elim": 19, "Death": 1, "Assists": 3, "Damage": 7762, "Heal": 0, "Mitigated": 3582, "Playtime": "08:15"},
    {"Team": "FRA", "Player": "XERION", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 13, "Death": 4, "Assists": 14, "Damage": 3371, "Heal": 4122, "Mitigated": 0, "Playtime": "08:15"},
    # AUS
    {"Team": "AUS", "Player": "ACKYYY", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 5, "Death": 8, "Assists": 5, "Damage": 1427, "Heal": 3145, "Mitigated": 0, "Playtime": "08:15"},
    {"Team": "AUS", "Player": "CUFFA", "Position": "TANK", "Map": "Samoa", "Map Type": "Control", "Elim": 6, "Death": 6, "Assists": 2, "Damage": 4156, "Heal": 0, "Mitigated": 5928, "Playtime": "08:15"},
    {"Team": "AUS", "Player": "COLOURHEX", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 9, "Death": 9, "Assists": 0, "Damage": 5963, "Heal": 0, "Mitigated": 169, "Playtime": "08:15"},
    {"Team": "AUS", "Player": "HAMSTER", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 4, "Death": 7, "Assists": 3, "Damage": 1970, "Heal": 5612, "Mitigated": 352, "Playtime": "08:15"},
    {"Team": "AUS", "Player": "SGY", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 5, "Death": 6, "Assists": 0, "Damage": 3016, "Heal": 0, "Mitigated": 0, "Playtime": "08:15"},

    # ==========================================
    # Match 3: FRA vs AUS (Map 2: Shambali Monastery)
    # ==========================================
    # FRA (Winner)
    {"Team": "FRA", "Player": "DIP", "Position": "DPS", "Map": "Shambali", "Map Type": "Escort", "Elim": 23, "Death": 6, "Assists": 0, "Damage": 7206, "Heal": 0, "Mitigated": 99, "Playtime": "11:45"},
    {"Team": "FRA", "Player": "FDGOD", "Position": "SPT", "Map": "Shambali", "Map Type": "Escort", "Elim": 24, "Death": 9, "Assists": 27, "Damage": 3214, "Heal": 8183, "Mitigated": 0, "Playtime": "11:45"},
    {"Team": "FRA", "Player": "KIO", "Position": "DPS", "Map": "Shambali", "Map Type": "Escort", "Elim": 18, "Death": 12, "Assists": 2, "Damage": 11096, "Heal": 0, "Mitigated": 295, "Playtime": "11:45"},
    {"Team": "FRA", "Player": "KROXZ", "Position": "TANK", "Map": "Shambali", "Map Type": "Escort", "Elim": 34, "Death": 6, "Assists": 5, "Damage": 16421, "Heal": 0, "Mitigated": 9297, "Playtime": "11:45"},
    {"Team": "FRA", "Player": "XERION", "Position": "SPT", "Map": "Shambali", "Map Type": "Escort", "Elim": 6, "Death": 9, "Assists": 17, "Damage": 4113, "Heal": 10406, "Mitigated": 1443, "Playtime": "11:45"},
    # AUS
    {"Team": "AUS", "Player": "ACKYYY", "Position": "SPT", "Map": "Shambali", "Map Type": "Escort", "Elim": 15, "Death": 11, "Assists": 19, "Damage": 3291, "Heal": 8646, "Mitigated": 0, "Playtime": "11:45"},
    {"Team": "AUS", "Player": "CUFFA", "Position": "TANK", "Map": "Shambali", "Map Type": "Escort", "Elim": 17, "Death": 8, "Assists": 3, "Damage": 10790, "Heal": 0, "Mitigated": 12334, "Playtime": "11:45"},
    {"Team": "AUS", "Player": "COLOURHEX", "Position": "DPS", "Map": "Shambali", "Map Type": "Escort", "Elim": 19, "Death": 12, "Assists": 4, "Damage": 11583, "Heal": 156, "Mitigated": 241, "Playtime": "11:45"},
    {"Team": "AUS", "Player": "HAMSTER", "Position": "SPT", "Map": "Shambali", "Map Type": "Escort", "Elim": 10, "Death": 14, "Assists": 9, "Damage": 4789, "Heal": 7978, "Mitigated": 869, "Playtime": "11:45"},
    {"Team": "AUS", "Player": "SGY", "Position": "DPS", "Map": "Shambali", "Map Type": "Escort", "Elim": 27, "Death": 5, "Assists": 0, "Damage": 10223, "Heal": 91, "Mitigated": 263, "Playtime": "11:45"},

    # ==========================================
    # Match 4: USA vs KOR (Map 1: Samoa)
    # ==========================================
    # KOR (Winner)
    {"Team": "KOR", "Player": "CHORONG", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 15, "Death": 3, "Assists": 18, "Damage": 3290, "Heal": 7137, "Mitigated": 0, "Playtime": "08:49"},
    {"Team": "KOR", "Player": "JUNBIN", "Position": "TANK", "Map": "Samoa", "Map Type": "Control", "Elim": 22, "Death": 3, "Assists": 10, "Damage": 10759, "Heal": 1806, "Mitigated": 5687, "Playtime": "08:49"},
    {"Team": "KOR", "Player": "SEONJUN", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 18, "Death": 6, "Assists": 3, "Damage": 8490, "Heal": 195, "Mitigated": 767, "Playtime": "08:49"},
    {"Team": "KOR", "Player": "SIMPLE", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 10, "Death": 4, "Assists": 12, "Damage": 5699, "Heal": 8420, "Mitigated": 0, "Playtime": "08:49"},
    {"Team": "KOR", "Player": "STALK3R", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 14, "Death": 5, "Assists": 2, "Damage": 8070, "Heal": 50, "Mitigated": 302, "Playtime": "08:49"},
    # USA
    {"Team": "USA", "Player": "HAWK", "Position": "TANK", "Map": "Samoa", "Map Type": "Control", "Elim": 11, "Death": 5, "Assists": 5, "Damage": 8442, "Heal": 2047, "Mitigated": 5538, "Playtime": "08:49"},
    {"Team": "USA", "Player": "LUKEMINO", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 12, "Death": 6, "Assists": 12, "Damage": 3362, "Heal": 12030, "Mitigated": 125, "Playtime": "08:49"},
    {"Team": "USA", "Player": "TR33", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 15, "Death": 6, "Assists": 1, "Damage": 10312, "Heal": 0, "Mitigated": 260, "Playtime": "08:49"},
    {"Team": "USA", "Player": "ULTRAVIOLET", "Position": "SPT", "Map": "Samoa", "Map Type": "Control", "Elim": 8, "Death": 7, "Assists": 8, "Damage": 5352, "Heal": 7624, "Mitigated": 1335, "Playtime": "08:49"},
    {"Team": "USA", "Player": "ZERUHH", "Position": "DPS", "Map": "Samoa", "Map Type": "Control", "Elim": 8, "Death": 4, "Assists": 1, "Damage": 4730, "Heal": 129, "Mitigated": 170, "Playtime": "08:49"},
]

df = pd.DataFrame(records)

# Save CSVs
csv_path1 = Path("clean_raw_owwc_day1.csv")
csv_path2 = Path("clean_raw_owwc_d1.csv")
df.to_csv(csv_path1, index=False, encoding="utf-8-sig")
df.to_csv(csv_path2, index=False, encoding="utf-8-sig")
print(f"✓ CSV saved: {csv_path1} ({len(df)} rows)")

# Save Excel files with premium styling
def save_styled_excel(df, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Day 1 Player Stats"

    # Header style
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Write headers
    headers = list(df.columns)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Team fills
    team_fills = {
        "KOR": PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid"),
        "KSA": PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid"),
        "FRA": PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid"),
        "ESP": PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid"),
        "AUS": PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid"),
        "USA": PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid"),
    }

    # Write data
    for r_idx, row in df.iterrows():
        row_num = r_idx + 2
        team = row["Team"]
        row_fill = team_fills.get(team, PatternFill(fill_type=None))

        for c_idx, col in enumerate(headers, 1):
            val = row[col]
            cell = ws.cell(row=row_num, column=c_idx, value=val)
            cell.font = bold_font if col in ["Player", "Team"] else data_font
            cell.fill = row_fill
            cell.border = border_thin

            if col in ["Elim", "Death", "Assists", "Damage", "Heal", "Mitigated"]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
            elif col in ["Team", "Position", "Map Type", "Playtime"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    wb.save(output_path)
    print(f"✓ Styled Excel saved: {output_path}")

save_styled_excel(df, Path("OWWC_2026_PLAYOFFS_DAY1.xlsx"))
save_styled_excel(df, Path("OWWC_PARSED_RESULT_D1.xlsx"))
print("All files generated successfully!")
