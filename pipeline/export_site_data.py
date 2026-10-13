"""Export a stage workbook to the Stat Lab site data format (window.OWCS_DATA shape).

Usage:
    python3 pipeline/export_site_data.py OWCS_STAT_LAB_26_KR_STAGE3.xlsx site/data/data_s3.js OWCS_DATA_S3

The workbook must have been saved by Excel (formula results cached).
Players with more than one row in TOTAL STATISTICS (position switch) are exported
as separate entries named "PLAYER (POS)", e.g. "LIP (DPS)" / "LIP (SPT)".
Map name / map type for player rows always come from MATCH INFO via MATCH_ID.
"""
import datetime as dt
import json
import sys
from collections import Counter, defaultdict

import openpyxl

WEEK_SHEETS = [f"MATCH_W{i}" for i in range(1, 7)]
STAT_KEYS = ["Elim", "Death", "Assists", "Damage", "Heal", "Mitigated"]


def ymd(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def hms(sec):
    sec = int(round(sec))
    return f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


def score(v):
    if v is None:
        return ""
    return str(int(v)) if float(v).is_integer() else f"{v:.2f}"


HERO_ALIASES = {"Rammatra": "Ramattra", "Lúcio": "Lucio"}
# T / D / S. Newer heroes were inferred from TEAMFIGHT INFO: every recorded comp is 1-2-2
# and only this assignment satisfies that. Add new heroes here (the export fails otherwise).
HERO_ROLES = {
    "D.Va": "T", "Doomfist": "T", "Mauga": "T", "Orisa": "T", "Ramattra": "T", "Reinhardt": "T",
    "Sigma": "T", "Winston": "T", "Wrecking Ball": "T", "D.Mon": "T", "Domina": "T",
    "Ashe": "D", "Bastion": "D", "Cassidy": "D", "Hanzo": "D", "Junkrat": "D", "Reaper": "D",
    "Sojourn": "D", "Sombra": "D", "Tracer": "D", "Widowmaker": "D", "Shion": "D", "Anran": "D",
    "Emre": "D", "Vendetta": "D",
    "Ana": "S", "Baptiste": "S", "Brigitte": "S", "Kiriko": "S", "Lucio": "S", "Zenyatta": "S",
    "Jetpack Cat": "S", "Wuyang": "S",
}



def hero(v):
    if not v:
        return ""
    name = str(v).strip().title()
    return HERO_ALIASES.get(name, name)


def rows(ws):
    header = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(header, r))


def read_teamfights(wb, mi):
    """Parse TEAMFIGHT INFO (Match_Id only on the first row of each map). Returns (teamfights, problems)."""
    teamfights, problems, cur_id, order = [], [], None, 0
    if "TEAMFIGHT INFO" not in wb.sheetnames:
        return teamfights, problems
    for r in wb["TEAMFIGHT INFO"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            cur_id, order = r[0], 0
        if not any(r[1:]):
            continue
        order += 1
        m = mi.get(cur_id)
        t1, t2 = [hero(h) for h in r[3:8]], [hero(h) for h in r[8:13]]
        where = f"TEAMFIGHT INFO {cur_id} #{order}"
        if not m:
            problems.append(f"{where}: MATCH_ID not in MATCH INFO")
            continue
        if r[13] not in (m["TEAM 1"], m["TEAM 2"]):
            problems.append(f"{where}: winner {r[13]} is not {m['TEAM 1']}/{m['TEAM 2']}")
        if len(set(filter(None, t1))) != 5 or len(set(filter(None, t2))) != 5:
            problems.append(f"{where}: each side needs 5 distinct heroes")
        problems += [f"{where}: no role for hero {h!r} (add it to HERO_ROLES)" for h in t1 + t2 if h not in HERO_ROLES]
        teamfights.append({"MATCH_ID": cur_id, "Order": order, "TEAM 1": m["TEAM 1"], "TEAM 2": m["TEAM 2"],
                           "TEAM 1 COMP": t1, "TEAM 2 COMP": t2, "WINNER": r[13]})
    return teamfights, problems


def main(xlsx, out_path, var_name):
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # --- totals / per10, detect position-split players ---
    totals = [r for r in rows(wb["TOTAL STATISTICS"]) if r["PLAYER"]]
    per10 = [r for r in rows(wb["Per10 STATISTICS"]) if r["Player"]]
    if any(r["TOTAL ELIMS"] is None for r in totals):
        sys.exit("Formula results are missing: open the workbook in Excel and save it first.")
    counts = Counter(r["PLAYER"] for r in totals)
    split = {p for p, n in counts.items() if n > 1}

    def key(player, pos):  # pos = map/roster position (TANK/DPS/SPT)
        return f"{player} ({pos})" if player in split else player

    # --- match info ---
    mi_rows = [r for r in rows(wb["MATCH INFO"]) if r["MAP"] and r["MATCH_ID"]]
    mi = {r["MATCH_ID"]: r for r in mi_rows}

    # --- raw player-map rows ---
    raw, problems = [], []
    for sheet in WEEK_SHEETS:
        for r in rows(wb[sheet]):
            if not r["Player"]:
                continue
            m = mi.get(r["MATCH_ID"])
            if not m:
                problems.append(f"{sheet}: MATCH_ID {r['MATCH_ID']} not in MATCH INFO")
                continue
            sec = int(round(r["Playtime(Seconds)"]))
            raw.append({
                "Team": r["Team"], "Player": key(r["Player"], r["Position"]), "Position": r["Position"],
                "Map": m["MAP"], "Map Type": m["MAP TYPE"],
                **{k: int(r[k]) for k in STAT_KEYS},
                "Playtime": hms(sec), "Playtime(Seconds)": sec, "MATCH_ID": r["MATCH_ID"],
            })

    per_map = Counter(r["MATCH_ID"] for r in raw)
    problems += [f"{mid}: {n} player rows (expected 10)" for mid, n in per_map.items() if n != 10]
    problems += [f"{mid}: no player rows" for mid in mi if mid not in per_map]
    map_len = defaultdict(int)
    for r in raw:
        map_len[r["MATCH_ID"]] = max(map_len[r["MATCH_ID"]], r["Playtime(Seconds)"])

    # the site groups maps into series by WEEK-DAY-MATCH, so each key must be a single pairing
    series = defaultdict(set)
    for r in mi_rows:
        series[(r["WEEK"], r["DAY"], r["MATCH"])].add((r["TEAM 1"], r["TEAM 2"]))
    problems += [f"MATCH INFO week {w} day {d} match {m} mixes series {sorted(v)}"
                 for (w, d, m), v in series.items() if len(v) > 1]

    match_info = []
    for r in mi_rows:
        s1, s2 = score(r["TEAM 1 SCORE"]), score(r["TEAM 2 SCORE"])
        unit = "m" if r["MAP TYPE"] == "Push" else ""
        match_info.append({
            "WEEK": r["WEEK"], "DAY": r["DAY"], "MATCH": r["MATCH"], "Set #": r["Set #"],
            "DATE": ymd(r["DATE"]), "PHASE": r["PHASE"], "TEAM 1": r["TEAM 1"], "TEAM 2": r["TEAM 2"],
            "MAP": r["MAP"], "MAP TYPE": r["MAP TYPE"], "TIME": hms(map_len[r["MATCH_ID"]]),
            "WINNER": r["WINNER"], "TEAM 1 BAN": hero(r["TEAM 1 BAN"]), "TEAM 2 BAN": hero(r["TEAM 2 BAN"]),
            "MATCH_ID": r["MATCH_ID"], "TEAM 1 SCORE": s1, "TEAM 2 SCORE": s2,
            "DETAIL_SCORE": f"{s1}{unit} : {s2}{unit}",
        })

    # --- team info (header row per team, then players) ---
    team_info, cur = [], None
    roster_pos = {}
    for r in rows(wb["TEAM INFO"]):
        if r["Team Name"]:
            cur = r["Shortened"]
            team_info.append({"Team Name": r["Team Name"], "Shortened": cur, "Player Name": None, "Player Position": None})
        if not r["Player Name"]:
            continue
        roster_pos[r["Player Name"]] = r["Player Position"]
        stints = [t for t in totals if t["PLAYER"] == r["Player Name"]] if r["Player Name"] in split else [None]
        for t in stints:
            pos = t["POSITION"] if t else r["Player Position"]
            team_info.append({"Team Name": None, "Shortened": None,
                              "Player Name": key(r["Player Name"], pos), "Player Position": pos,
                              "Detailed": (t["DETAILED"] if t else r["Detailed"]) or ""})

    # --- totals / per10 in site shape (Position = detailed role) ---
    total_stats = [{
        "TEAM": t["TEAM"], "PLAYER": key(t["PLAYER"], t["POSITION"]), "POSITION": t["DETAILED"] or t["POSITION"],
        "TOTAL ELIMS": t["TOTAL ELIMS"], "TOTAL DEATHS": t["TOTAL DEATHS"], "TOTAL ASSISTS": t["TOTAL ASSISTS"],
        "TOTAL DAMAGE": t["TOTAL DAMAGE"], "TOTAL HEAL": t["TOTAL HEAL"], "TOTAL MITIGATED": t["TOTAL MITIGATED"],
        "TOTAL E/D": round(t["TOTAL E/D"], 3), "TOTAL PLAYTIME": round(t["TOTAL PLAYTIME"]),
    } for t in totals]
    per10_stats = [{
        "Team": p["Team"], "Player": key(p["Player"], p["Position"]), "Position": p["DETAILED"] or p["Position"],
        **{k: round(p[k], 2) for k in ["Playtime_Min", "Elim / 10", "Death / 10", "Assists / 10",
                                         "Damage / 10", "Heal / 10", "Mitigated / 10"]},
    } for p in per10]

    # cross-check workbook totals against the raw rows we export
    agg = defaultdict(lambda: [0] * 7)
    for r in raw:
        a = agg[r["Player"]]
        for i, k in enumerate(STAT_KEYS):
            a[i] += r[k]
        a[6] += r["Playtime(Seconds)"]
    for t in total_stats:
        expect = [t[k] for k in ["TOTAL ELIMS", "TOTAL DEATHS", "TOTAL ASSISTS", "TOTAL DAMAGE",
                                 "TOTAL HEAL", "TOTAL MITIGATED"]] + [t["TOTAL PLAYTIME"]]
        got = agg.pop(t["PLAYER"], [0] * 7)
        if [round(x) for x in expect] != [round(x) for x in got]:
            problems.append(f"totals mismatch for {t['PLAYER']}: workbook {expect} vs raw {got}")
    problems += [f"raw rows for {p} have no TOTAL STATISTICS row" for p in agg]

    teamfights, tf_problems = read_teamfights(wb, mi)
    problems += tf_problems

    potm = [{"Date": ymd(r["Date"]), "Match": r["Match"],
             "POTM": key(r["POTM"], r["Position"]), "Position": r["Position"]}
            for r in rows(wb["POTM"]) if r["POTM"]]

    data = {"teamInfo": team_info, "matchInfo": match_info, "totalStats": total_stats,
            "per10Stats": per10_stats, "potm": potm, "rawStats": raw, "teamfights": teamfights,
            "heroRoles": HERO_ROLES if teamfights else {}}
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(f"window.{var_name} = {json.dumps(data, ensure_ascii=False)};\n")

    print(f"{out_path}: {len(match_info)} maps, {len(raw)} player-map rows, "
          f"{len(total_stats)} players, {len(potm)} POTM, "
          f"{len(teamfights)} teamfights on {len({t['MATCH_ID'] for t in teamfights})} maps; split players: {sorted(split) or 'none'}")
    if problems:
        print("PROBLEMS:\n  " + "\n  ".join(problems))
        sys.exit(1)


if __name__ == "__main__":
    main(*sys.argv[1:4])
