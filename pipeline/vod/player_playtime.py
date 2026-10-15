"""Attach players to HUD slots and compare VOD playtime with the official per-map playtime.

    python pipeline/vod/player_playtime.py OWCS_STAT_LAB_26_KR_STAGE3.xlsx \
        tmp/vod/plates/plate_readings.csv tmp/vod/s3w1_hero_playtime.csv tmp/vod/s3w1_player_hero_playtime.csv

plate_readings.csv: per map and HUD side, the five name plates left to right (read from 720p
frames). Each name is mapped to the canonical TEAM INFO player only through display-font
confusables (0/O, 1/I); anything else stops the run. The five players per team-map must equal
that map's MATCH_W1 players, otherwise the run stops (identity conflicts are not resolved here).
Output adds player and official_seconds (MATCH_W1 Playtime(Seconds)) to every hero row and
prints per-map coverage = VOD counted seconds / official seconds.
"""
import csv
import sys

import openpyxl
import pandas as pd


def fold(s):
    return str(s).upper().replace("0", "O").replace("1", "I")


def main(xlsx, plates_csv, hero_csv, out_csv):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    team, roster = None, {}
    for r in list(wb["TEAM INFO"].iter_rows(values_only=True))[1:]:
        team = r[1] or team
        if r[2]:
            roster.setdefault(team, {})[fold(r[2])] = r[2]
    head, *body = list(wb["MATCH_W1"].iter_rows(values_only=True))
    mw = pd.DataFrame([r for r in body if r[head.index("MATCH_ID")]], columns=head)
    official = mw.set_index(["MATCH_ID", "Player"])["Playtime(Seconds)"].to_dict()
    truth = mw.groupby(["MATCH_ID", "Team"])["Player"].apply(set).to_dict()

    slot_player, errors = {}, []
    for r in csv.DictReader(open(plates_csv, encoding="utf-8")):
        names = [r[f"p{i}"] for i in range(1, 6)]
        mapped = [roster[r["team_on_hud"]].get(fold(n)) for n in names]
        if None in mapped:
            errors.append(f"{r['match_id']} {r['team_on_hud']}: not in roster {names}")
            continue
        if set(mapped) != truth.get((r["match_id"], r["team_on_hud"])):
            errors.append(f"{r['match_id']} {r['team_on_hud']}: plates {sorted(mapped)} != MATCH_W1")
        base = 0 if r["hud_side"] == "left" else 5
        for i, p in enumerate(mapped):
            slot_player[(r["match_id"], base + i)] = (r["team_on_hud"], p)
    if errors:
        sys.exit("Identity problems:\n  " + "\n  ".join(errors))

    d = pd.read_csv(hero_csv, encoding="utf-8-sig")
    d["plate_team"] = [slot_player[(m, s)][0] for m, s in zip(d.match_id, d.slot)]
    d["player"] = [slot_player[(m, s)][1] for m, s in zip(d.match_id, d.slot)]
    bad = d[d.plate_team != d.team]
    if len(bad):
        sys.exit(f"Side assignment disagrees with name plates on maps: {sorted(set(bad.match_id))}")
    d["official_seconds"] = [official.get((m, p)) for m, p in zip(d.match_id, d.player)]
    d.drop(columns="plate_team").to_csv(out_csv, index=False, encoding="utf-8-sig")
    per = d.groupby(["match_id", "player"]).agg(vod=("seconds", "sum"), off=("official_seconds", "first"))
    per["cov"] = per.vod / per.off
    print(f"{out_csv}: {len(d)} rows, {per.index.size} player-maps")
    print(f"coverage VOD/official: median {per['cov'].median():.1%}, "
          f"p10 {per['cov'].quantile(.1):.1%}, p90 {per['cov'].quantile(.9):.1%}, max {per['cov'].max():.1%}")
    m = per.groupby(level=0)["cov"].median().sort_values()
    print("lowest maps:", ", ".join(f"{k} {v:.0%}" for k, v in m.head(5).items()))


if __name__ == "__main__":
    main(*sys.argv[1:])
