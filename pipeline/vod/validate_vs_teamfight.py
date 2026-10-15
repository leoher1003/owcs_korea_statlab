"""Compare VOD hero classification with the manually recorded TEAMFIGHT INFO (read-only).

    python pipeline/vod/validate_vs_teamfight.py OWCS_STAT_LAB_26_KR_STAGE3.xlsx tmp/vod/261002_scan.npz ...

Steps
  1. Map segments per broadcast day: runs of in-game HUD frames, joining gaps shorter than
     GAP_S (replays, camera cuts), keeping runs of at least MIN_MAP_S.
  2. Segments are paired in order with that day's MATCH INFO rows (match, then set order).
     A segment count that differs from the MATCH INFO count stops the day; nothing is guessed.
  3. Per map and side: heroes the classifier saw for at least MIN_HERO_S (NONE / UNKNOWN /
     NO_HUD ignored) vs the union of heroes in that map's TEAMFIGHT INFO compositions.
     HUD slots 0-4 are the left team; it is compared as MATCH INFO team 1, and the swapped
     orientation is also scored so a side flip is visible instead of silently "fixed".
Only maps with a teamfight log are scored. TEAMFIGHT INFO compositions are start-of-fight
snapshots, so a hero played briefly between fights can be missing from the log; extra VOD
heroes are listed for review, not counted as classifier errors outright.
"""
import sys
from pathlib import Path

import numpy as np
import openpyxl

GAP_S = 180
MIN_MAP_S = 180
MIN_HERO_S = 15
FPS = 2


def segments(on):
    idx = np.flatnonzero(on)
    br = np.flatnonzero(np.diff(idx) > GAP_S * FPS)
    st, en = np.r_[idx[0], idx[br + 1]], np.r_[idx[br], idx[-1]]
    return [(a, b) for a, b in zip(st, en) if (b - a) / FPS >= MIN_MAP_S]


def read_master(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    mi = list(wb["MATCH INFO"].iter_rows(values_only=True))
    head = mi[0]
    maps = [dict(zip(head, r)) for r in mi[1:] if r[head.index("MATCH_ID")] and "#" not in str(r[head.index("MATCH_ID")])]
    fights, mid = {}, None
    for r in list(wb["TEAMFIGHT INFO"].iter_rows(values_only=True))[1:]:
        mid = r[0] or mid
        if mid and any(r[3:13]):
            fights.setdefault(mid, []).append(([str(h).strip() for h in r[3:8] if h], [str(h).strip() for h in r[8:13] if h]))
    return maps, fights


def vod_heroes(hero, names, a, b, slots):
    h = hero[a:b + 1, slots].ravel()
    skip = {names.index(n) for n in ("NONE", "UNKNOWN", "NO_HUD")}
    ids, cnt = np.unique(h, return_counts=True)
    return {names[i].upper(): c / FPS for i, c in zip(ids, cnt) if i not in skip and c / FPS >= MIN_HERO_S}


def main(xlsx, *scans):
    maps, fights = read_master(xlsx)
    totals = {"rec": 0, "hit": 0, "extra": 0, "flip": 0, "scored": 0}
    for s in scans:
        day = Path(s).name[:6]
        z = np.load(Path(s).with_name(f"{day}_heroes.npz"))
        hero, names = z["hero"], [str(n) for n in z["names"]]
        seg = segments(hero[:, 0] != names.index("NO_HUD"))
        rows = [m for m in maps if str(m["MATCH_ID"]).startswith(day)]
        print(f"\n== {day}: {len(seg)} HUD segments, {len(rows)} MATCH INFO maps")
        if len(seg) != len(rows):
            print("   segment count mismatch; skipping this day")
            continue
        for (a, b), m in zip(seg, rows):
            mid = m["MATCH_ID"]
            left, right = vod_heroes(hero, names, a, b, range(5)), vod_heroes(hero, names, a, b, range(5, 10))
            if mid not in fights:
                print(f"   {mid:22s} {a / 120:5.0f}-{b / 120:3.0f} min  no teamfight log")
                continue
            t1 = set().union(*[set(f[0]) for f in fights[mid]])
            t2 = set().union(*[set(f[1]) for f in fights[mid]])
            straight = len(t1 & set(left)) + len(t2 & set(right))
            swapped = len(t1 & set(right)) + len(t2 & set(left))
            flip = swapped > straight
            v1, v2 = (right, left) if flip else (left, right)
            rec, hit = len(t1) + len(t2), max(straight, swapped)
            extra = (set(v1) - t1) | (set(v2) - t2)
            miss = (t1 - set(v1)) | (t2 - set(v2))
            totals["rec"] += rec
            totals["hit"] += hit
            totals["extra"] += len(extra)
            totals["flip"] += flip
            totals["scored"] += 1
            print(f"   {mid:22s} {a / 120:5.0f}-{b / 120:3.0f} min  recall {hit}/{rec}"
                  f"{'  SIDE FLIPPED' if flip else ''}"
                  f"{'  missing ' + ','.join(sorted(miss)) if miss else ''}"
                  f"{'  extra ' + ','.join(f'{h}({(v1 | v2).get(h, 0):.0f}s)' for h in sorted(extra)) if extra else ''}")
    t = totals
    if t["scored"]:
        print(f"\nScored {t['scored']} maps: recorded heroes found {t['hit']}/{t['rec']} ({t['hit'] / t['rec']:.1%}), "
              f"extra VOD heroes {t['extra']}, side flips {t['flip']}")


if __name__ == "__main__":
    main(*sys.argv[1:])
